"""Turns a validated payload into the bytes of an .xlsx file, with openpyxl.

Nothing here is styled. The only thing a cell can carry besides its value is a
number format, and that is one of the strings in contract.json.
"""
from __future__ import annotations

import io

from openpyxl import Workbook as OpenpyxlWorkbook
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.utils.exceptions import IllegalCharacterError

from .contract import CREATOR, ContractError, Moment, Number, Sheet, Text, Workbook


def build_xlsx(workbook: Workbook) -> bytes:
    book = OpenpyxlWorkbook()
    # Set rather than left to the library's default so the signature is a line of
    # this file and a test, not an accident of a dependency's next release.
    book.properties.creator = CREATOR
    for index, sheet in enumerate(workbook.sheets):
        target = book.active if index == 0 else book.create_sheet()
        target.title = sheet.name
        _write_sheet(target, sheet)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _write_sheet(target, sheet: Sheet) -> None:
    for row_number, row in enumerate(sheet.rows, 1):
        for column_number, cell in enumerate(row, 1):
            if cell is None:
                continue
            _write_cell(target.cell(row=row_number, column=column_number), cell, sheet.name)


def _write_cell(target: OpenpyxlCell, cell: Text | Number | Moment, sheet: str) -> None:
    try:
        if isinstance(cell, Text):
            _write_text(target, cell.value)
        else:
            target.value = cell.value
            if cell.fmt:
                target.number_format = cell.fmt
    except IllegalCharacterError as error:
        raise ContractError(
            "illegal_character",
            "a control character Excel cannot store",
            sheet=sheet,
            row=target.row,
            column=target.column,
        ) from error


def _write_text(target: OpenpyxlCell, value: str) -> None:
    # openpyxl decides a string that starts with "=" is a formula and would write
    # it as one. Assigning the type afterwards is what makes every string in the
    # payload a string, whatever it starts with: a Beleg description of
    # =HYPERLINK(...) is data, not an instruction.
    target.value = value
    target.data_type = "s"
