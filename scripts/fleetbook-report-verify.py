"""Read both files of a simulated month back and check them, from their own bytes.

    python scripts/fleetbook-report-verify.py <SIMULIERT_Monatsbericht_….xlsx> <SIMULIERT_Betriebsbericht_….xlsx>

External audit round 6. The month comes as two files: the Fahrtenbuch report
(passenger rides and cancellations) and the internal operational report (empty
drives, declined orders, scheduling). This script imports nothing from the
generator and nothing from JavaScript: it opens both files with openpyxl and
recomputes, in integer cents and on absolute instants (Europe/Berlin through
zoneinfo):

  - the label on every data row of every sheet, and no formula anywhere;
  - zero declined orders in the Fahrtenbuch report, in any sheet, by status
    and by id, and every one of them in the operational report;
  - the money of the report's own scope: fares, tips and cancellation fees of
    A re-summed against B, C, D; the cancellation fee per row against the
    company setting in I; the basis, amount, remainder and rate of F;
  - the calendar: every day of the month a row of B, 24 h or 23/25;
  - overlaps: no vehicle and no driver (by id) holds two rides at once;
  - shift lengths 240 to 360 minutes, from H2 and from O;
  - TÜV and confirmed full-day blocks, read from H and H2;
  - the route sheet G: one row per consecutive completed pair, its gap
    recomputed from A, and for a verified pair road + 45 s + margin <= gap;
  - every empty drive of L: it names a completed ride of A with the same
    vehicle, driver and shift, its margin is recomputed from the options in I,
    and a verified one fits its window;
  - the verdict in F is not VALIDATED while a pair is estimated, unverifiable
    or in conflict.

One line per check; a check it cannot make says so and is never counted as
passed. Exit 1 when any check fails.
"""
import calendar
import json
import math
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

BERLIN = ZoneInfo("Europe/Berlin")
UTC = ZoneInfo("UTC")
REPORT_SHEETS = [
    "A Fahrten", "B Monatsuebersicht", "C Summe je Fahrer", "D Summe je Fahrzeug",
    "E Stornierungen", "F Pruefung und Grenzen", "G Routenpruefung",
    "H Ereignisse", "H2 Schichtentscheidungen", "I Quelle und Einstellungen", "SIMULIERT",
]
OPERATIONAL_SHEETS = [
    "L Leerfahrten", "M Abgelehnt intern", "N Interne Einnahmen", "O Schichtplanung",
    "I Quelle und Einstellungen", "SIMULIERT",
]
STATUSES = {"ROUTE_VERIFIED", "ROUTE_CONFLICT", "UNVERIFIABLE", "ESTIMATE_ONLY"}
HANDLING = 45


def cents(x):
    return round(float(x) * 100) if x not in (None, "") else 0


def instant(text):
    """'2026-10-25 02:30:00.123' on the Berlin clock, as an absolute UTC instant; None when not a time."""
    if not isinstance(text, str) or len(text) < 19:
        return None
    wall = datetime.strptime(text[:19], "%Y-%m-%d %H:%M:%S")
    millis = int(text[20:23]) if len(text) >= 23 and text[19] == "." else 0
    local = wall.replace(tzinfo=BERLIN, fold=0)
    back = local.astimezone(UTC).astimezone(BERLIN).replace(tzinfo=None)
    if back != wall:
        return None
    return local.astimezone(UTC) + timedelta(milliseconds=millis)


def table(book, name):
    rows = list(book[name].iter_rows(values_only=True))
    header = list(rows[0])
    return header, [dict(zip(header, r)) for r in rows[1:]]


class Checks:
    def __init__(self):
        self.lines, self.failures = [], []

    def __call__(self, ok, what, detail=""):
        self.lines.append(f"{'ok  ' if ok else 'FAIL'} {what}{'  ' + detail if detail else ''}")
        if not ok:
            self.failures.append(what)

    def info(self, text):
        self.lines.append(f"info {text}")


def labels(check, book, sheets, which):
    for name in sheets[:-1]:
        rows = list(book[name].iter_rows(values_only=True))
        bad = [i + 2 for i, r in enumerate(rows[1:]) if r[0] != "SIMULIERT"]
        check(rows[0][0] == "Kennzeichnung" and not bad, f"{which}: label on every row of {name}", f"{len(rows) - 1} rows")
    formulas = sum(1 for ws in book.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f")
    check(formulas == 0, f"{which}: no formula anywhere", f"{formulas} formula cells")


def settings_of(book):
    _, rows = table(book, "I Quelle und Einstellungen")
    return {r["Angabe"]: r["Wert"] for r in rows}


def verify(report_path, ops_path):
    check = Checks()
    book = load_workbook(report_path, read_only=False)
    ops = load_workbook(ops_path, read_only=False)
    check(book.sheetnames == REPORT_SHEETS, "report: sheets and their order", " | ".join(book.sheetnames))
    check(ops.sheetnames == OPERATIONAL_SHEETS, "operational: sheets and their order", " | ".join(ops.sheetnames))
    labels(check, book, REPORT_SHEETS, "report")
    labels(check, ops, OPERATIONAL_SHEETS, "operational")
    source = settings_of(book)
    check(source == settings_of(ops), "both files carry the same source sheet (run id, seed, settings)",
          f"run {source.get('Lauf-ID')}")

    # ── A: the report's rides ────────────────────────────────────────────────
    _, rides = table(book, "A Fahrten")
    by_id = {r["Fahrt-ID"]: r for r in rides}
    done = [r for r in rides if r["Status"] == "abgeschlossen"]
    cancelled = [r for r in rides if r["Status"] == "storniert"]
    check([r["Nr"] for r in rides] == list(range(1, len(rides) + 1)), "A: row numbers 1..n without gap", f"{len(rides)} rides")
    check(len(by_id) == len(rides), "A: every Fahrt-ID unique")
    check({r["Status"] for r in rides} <= {"abgeschlossen", "storniert"}, "A: only completed and cancelled rides",
          json.dumps(Counter(r["Status"] for r in rides)))

    _, declined = table(ops, "M Abgelehnt intern")
    declined_ids = {r["Fahrt-ID"] for r in declined}
    leaks = []
    for ws in book.worksheets:
        for row in ws.iter_rows(values_only=True):
            if "abgelehnt" in row or any(isinstance(c, str) and c in declined_ids for c in row):
                leaks.append(ws.title)
    check(not leaks, "report: zero declined orders, by status and by id, in every sheet",
          f"{len(declined)} declined kept in M; leaks in {sorted(set(leaks))}")

    fare = sum(cents(r["Fahrpreis (EUR)"]) for r in done)
    tips = sum(cents(r["Trinkgeld (EUR)"]) for r in done)
    storno = sum(cents(r["Stornogebuehr (EUR)"]) for r in cancelled)
    charged = source.get("Einstellung cancellationFeeEnabled") == "Ja"
    unit = int(source.get("Einstellung cancellationFeeCents") or 0)
    expected_fee = unit if charged else 0
    check(all(cents(r["Stornogebuehr (EUR)"]) == expected_fee and cents(r["Fahrpreis (EUR)"]) == 0 for r in cancelled),
          "A: every cancellation carries the company fee and no fare", f"{len(cancelled)} x {expected_fee} cents")
    check(all(r["Finanzstatus"] for r in rides), "A: every ride states its financial status")
    check(all(cents(r["Stornogebuehr (EUR)"]) == 0 for r in done), "A: no completed ride carries a cancellation fee")

    for name, key in (("B Monatsuebersicht", "Datum"), ("C Summe je Fahrer", "Fahrer"), ("D Summe je Fahrzeug", "Fahrzeug")):
        _, rows = table(book, name)
        total = [r for r in rows if r[key] == "Summe"]
        parts = [r for r in rows if r[key] != "Summe"]
        ok = len(total) == 1
        for col, want in (("Fahrpreise (EUR)", fare), ("Trinkgeld (EUR)", tips), ("Stornogebuehren (EUR)", storno)):
            check(ok and cents(total[0][col]) == want == sum(cents(r[col]) for r in parts), f"{col} of A = Summe of {name} = its parts",
                  f"{want / 100:.2f} EUR")
        check(ok and total[0]["Fahrten"] == len(rides) == sum(r["Fahrten"] for r in parts), f"rides of A = Summe of {name} = its parts")
    _, e_rows = table(book, "E Stornierungen")
    check(len(e_rows) == len(cancelled) and sum(cents(r["Stornogebuehr (EUR)"]) for r in e_rows) == storno,
          "cancellations of A = sheet E", f"{len(e_rows)} rows, {storno / 100:.2f} EUR")

    # ── F: basis, share, verdict ─────────────────────────────────────────────
    _, f_rows = table(book, "F Pruefung und Grenzen")
    posten = {r["Posten"]: r["Wert / Summe der Teile"] for r in f_rows if r["Abschnitt"] == "Einnahmengrundlage und Monatsanteil"}
    verify_rows = {r["Posten"]: r for r in f_rows if r["Abschnitt"] == "Pruefung fuer den Export"}
    check(all(r["Ergebnis"] == "OK" and r["Differenz"] == 0 for r in f_rows if r["Abschnitt"] == "Abgleich"),
          "F: every reconciliation line is OK")
    basis = fare + (tips if posten.get("Trinkgeld in Grundlage") == "Ja" else 0)
    basis += storno if posten.get("Stornogebuehren in Grundlage") == "Ja" else 0
    check(cents(posten.get("Anspruchsberechtigte Einnahmen (EUR)")) == basis, "F: basis re-derived from A, no decline fee in it",
          f"{basis / 100:.2f} EUR")
    if "Satz (%)" in posten:
        amount, rest = cents(posten["Betrag (EUR)"]), cents(posten["Rest, falls Abzug (EUR)"])
        rate = float(posten["Satz (%)"])
        bp = round(rate * 100)
        check(amount + rest == basis, "F: amount + remainder = basis", f"{amount / 100:.2f} + {rest / 100:.2f}")
        check(30.0 <= rate <= 33.0 and amount == (basis * bp + 5000) // 10000, "F: rate in [30 %, 33 %] and amount = basis x rate",
              f"{rate} %")
    else:
        check(False, "F: the monthly share is in the file", "no 'Satz (%)' line")
    verdict = verify_rows.get("Exportstatus", {}).get("Wert / Summe der Teile")
    check(verdict in ("VALIDATED_FOR_EXPORT", "VERIFICATION_INCOMPLETE", "NOT_VALIDATED"), "F: states a verdict", str(verdict))
    check(str(verdict) in str(report_path), "the verdict is in the file name", str(verdict))

    # ── B: the calendar ──────────────────────────────────────────────────────
    _, b_rows = table(book, "B Monatsuebersicht")
    days = [r for r in b_rows if r["Datum"] != "Summe"]
    first = days[0]["Datum"]
    year, month = int(first[:4]), int(first[5:7])
    expected_days = [f"{year}-{month:02d}-{d:02d}" for d in range(1, calendar.monthrange(year, month)[1] + 1)]
    check([r["Datum"] for r in days] == expected_days, "B: every calendar day is a row", f"{len(days)} days")
    odd = {r["Datum"]: r["Stunden"] for r in days if r["Stunden"] != 24}
    check(all(h in (23, 25) for h in odd.values()), "B: a day is 24 h, or 23/25 on a clock change", json.dumps(odd))
    zero = [r for r in days if r["Tag ohne Fahrt"] == "Ja"]
    check(all(r["Fahrten"] == 0 and r["Grund"] for r in zero), "B: every zero-ride day has no ride and a reason", f"{len(zero)} days")

    # ── overlaps and shifts ──────────────────────────────────────────────────
    unreadable = [r["Fahrt-ID"] for r in done if not (instant(r["Fahrtbeginn"]) and instant(r["Fahrtende"]))]
    check(not unreadable, "A: every completed ride has a readable start and end", f"{len(unreadable)} unreadable")

    def overlaps(key):
        groups = defaultdict(list)
        for r in done:
            s, e = instant(r["Fahrtbeginn"]), instant(r["Fahrtende"])
            if s and e:
                groups[r[key]].append((s, e))
        return sum(sum(1 for a, b in zip(sorted(v), sorted(v)[1:]) if b[0] < a[1]) for v in groups.values())
    check(overlaps("Kennzeichen") == 0, "no vehicle holds two rides at once")
    check(overlaps("Fahrer-ID") == 0, "no driver (by id) holds two rides at once")

    _, decisions = table(book, "H2 Schichtentscheidungen")
    accepted = [d for d in decisions if str(d["Ergebnis"]).startswith("eingeplant")]
    lengths = [(instant(d["Schichtende"] + ".000") - instant(d["Schichtbeginn"] + ".000")).total_seconds() / 60 for d in accepted]
    outside = [round(m) for m in lengths if m < 240 or m > 360]
    check(accepted and not outside, "H2: every accepted shift is 240 to 360 minutes",
          f"{len(accepted)} shifts, min {round(min(lengths))} max {round(max(lengths))}" if lengths else "no shift")
    _, planning = table(ops, "O Schichtplanung")
    planned = [r for r in planning if r["Art"] == "Schicht eingeplant"]
    check(len(planned) == len(accepted) and all(240 <= r["Dauer (min)"] <= 360 for r in planned),
          "O: the same shifts, all within their limits", f"{len(planned)} shifts")
    short = sum(r["Anzahl"] for r in planning if r["Art"] == "Versuch verworfen" and "Mindestdauer" in str(r["Grund"]))
    check.info(f"{short} shift attempts refused at scheduling time for ending short of 240 minutes")

    blocked_days = {(r["Kennzeichen"], str(r["Tag"])) for r in decisions if r.get("Code") == "vehicle_tuev_expired"}
    on_blocked = [r["Fahrt-ID"] for r in done if (r["Kennzeichen"], str(r["Datum der Fahrt"])) in blocked_days]
    check(not on_blocked, "no ride of a vehicle on a day its TÜV had expired", f"{len(blocked_days)} vehicle-days blocked")

    _, events = table(book, "H Ereignisse")

    def window(text):
        parts = str(text or "").split(" bis ")
        if len(parts) != 2 or len(parts[0]) != 10 or len(parts[1]) != 10:
            return None
        a = datetime.strptime(parts[0], "%Y-%m-%d").replace(tzinfo=BERLIN)
        b = (datetime.strptime(parts[1], "%Y-%m-%d") + timedelta(days=1)).replace(tzinfo=BERLIN)
        return a, b
    confirmed = [e for e in events if e.get("Betrifft") == "Fahrer" and e.get("Quelle") == "driver_note"
                 and e.get("Sicherheit") == "bestaetigt" and e.get("Wirkung") == "hard_block"]
    inside = []
    for e in confirmed:
        w = window(e.get("Zeitraum"))
        if w is None:
            check(False, "a confirmed block with clock times", "not checked by this script")
            continue
        for r in done:
            s_, e_ = instant(r["Fahrtbeginn"]), instant(r["Fahrtende"])
            if r["Fahrer"] == e.get("Bezeichnung") and s_ and e_ and s_ < w[1] and e_ > w[0]:
                inside.append(r["Fahrt-ID"])
    check(not inside, "no ride of a driver inside a confirmed full-day block", f"{len(confirmed)} blocks")
    review = Counter(d["Code"] for d in decisions if d["Ergebnis"] == "eingeplant, zu pruefen")
    check.info(f"shifts accepted for human review, by cause: {json.dumps(review)}")

    # ── G: the route check ───────────────────────────────────────────────────
    margin_s = int(source.get("Option repositioningMarginSeconds") or 0)
    margin_p = float(source.get("Option repositioningMarginPercent") or 0)

    def margin(road):
        return margin_s + math.ceil(road * margin_p / 100)

    _, g_rows = table(book, "G Routenpruefung")
    pairs = set()
    for key in ("Kennzeichen", "Fahrer-ID"):
        groups = defaultdict(list)
        for r in done:
            groups[r[key]].append((instant(r["Fahrtbeginn"]), r["Fahrt-ID"]))
        for spans in groups.values():
            spans.sort()
            pairs |= {(key, a[1], b[1]) for a, b in zip(spans, spans[1:])}
    check(len(g_rows) == len(pairs), "G: one row per consecutive completed pair (vehicle and driver)",
          f"{len(g_rows)} rows, {len(pairs)} pairs")
    status = Counter(r["Status"] for r in g_rows)
    check(set(status) <= STATUSES, "G: only the four statuses", json.dumps(status))
    gap_wrong, fit_wrong = [], []
    for r in g_rows:
        a, b = by_id.get(r["Fahrt-ID frueher"]), by_id.get(r["Fahrt-ID spaeter"])
        if not (a and b):
            gap_wrong.append(r["Fahrt-ID frueher"])
            continue
        gap = round((instant(b["Fahrtbeginn"]) - instant(a["Fahrtende"])).total_seconds())
        if gap != r["Luecke (s)"]:
            gap_wrong.append(r["Fahrt-ID frueher"])
        if r["Status"] == "ROUTE_VERIFIED":
            need = r["Fahrzeit Strasse (s)"] + HANDLING + margin(r["Fahrzeit Strasse (s)"])
            if need != r["Benoetigt (s)"] or gap < need:
                fit_wrong.append(r["Fahrt-ID frueher"])
    check(not gap_wrong, "G: every gap recomputed from A's own times", f"{len(gap_wrong)} differ")
    check(not fit_wrong, "G: every verified pair holds road + 45 s + margin", f"{len(fit_wrong)} do not")
    same_shift_bad = [r for r in g_rows if r["Gleiche Schicht"] == "Ja" and r["Status"] in ("ROUTE_CONFLICT", "UNVERIFIABLE")]
    check(not same_shift_bad, "G: no conflict and no unverifiable pair inside a shift", f"{len(same_shift_bad)}")
    open_pairs = status.get("ROUTE_CONFLICT", 0) + status.get("UNVERIFIABLE", 0) + status.get("ESTIMATE_ONLY", 0)
    check(not (open_pairs and verdict == "VALIDATED_FOR_EXPORT"), "F: not VALIDATED while a pair is open",
          f"{open_pairs} open pairs, verdict {verdict}")

    # ── L: the empty drives ──────────────────────────────────────────────────
    _, legs = table(ops, "L Leerfahrten")
    leg_status = Counter(r["Pruefstatus"] for r in legs)
    check(set(leg_status) <= STATUSES, "L: only the four statuses", json.dumps(leg_status))
    check({(r["Art"], r["Bezeichnung"], r["Datenart"]) for r in legs} == {("REPOSITIONING", "LEERFAHRT", "SIMULATED")},
          "L: every record is a simulated REPOSITIONING / LEERFAHRT", f"{len(legs)} legs")
    check({r["Lauf-ID"] for r in legs} == {source.get("Lauf-ID")}, "L: one run id, the source sheet's")
    orphan, owner_wrong, leg_bad = [], [], []
    for leg in legs:
        nxt = by_id.get(leg["Fahrt-ID nachher"])
        if not nxt or nxt["Status"] != "abgeschlossen":
            orphan.append(leg["Leerfahrt-ID"])
            continue
        if (nxt["Kennzeichen"], nxt["Fahrer-ID"], nxt["Schicht"]) != (leg["Kennzeichen"], leg["Fahrer-ID"], leg["Schicht"]):
            owner_wrong.append(leg["Leerfahrt-ID"])
        if leg["Zweck"] != "Zielort zum Warteplatz" and leg["Nach"] != nxt["Abholort"]:
            leg_bad.append(leg["Leerfahrt-ID"])
        prev = by_id.get(leg["Fahrt-ID vorher"]) if leg["Fahrt-ID vorher"] else None
        if leg["Zweck"] in ("Zielort zum Warteplatz",) and (not prev or leg["Von"] != prev["Zielort"]):
            leg_bad.append(leg["Leerfahrt-ID"])
        if leg["Pruefstatus"] == "ROUTE_VERIFIED":
            road = leg["Fahrzeit Strasse (s)"]
            handling = 0 if leg["Zweck"] == "Zielort zum Warteplatz" else HANDLING
            need = road + handling + margin(road)
            arrive = instant(leg["Ankunft"])
            depart = instant(leg["Abfahrt"])
            if need != leg["Benoetigt (s)"] or leg["Verfuegbar (s)"] < need or not (arrive and depart) \
                    or round((arrive - depart).total_seconds()) != road:
                leg_bad.append(leg["Leerfahrt-ID"])
    check(not orphan, "L: every leg ends at a completed ride of A", f"{len(orphan)} do not")
    check(not owner_wrong, "L: every leg carries its ride's vehicle, driver and shift", f"{len(owner_wrong)} differ")
    check(not leg_bad, "L: every verified leg's places, arrival and margin re-derive and fit", f"{len(set(leg_bad))} do not")
    with_road = [r for r in legs if r["Fahrzeit Strasse (s)"] not in (None, "")]
    preferred = [r for r in with_road if r["Im Wunschfenster 2-5 min"] == "Ja"]
    check.info(f"{len(preferred)} of {len(with_road)} legs with a route are 2-5 minutes; "
               f"longest {max((r['Fahrzeit Strasse (s)'] for r in with_road), default=0)} s; "
               f"{sum(float(r['Strecke (km)'] or 0) for r in with_road):.1f} empty km")

    # ── M, N: declined orders, internal ──────────────────────────────────────
    fee_unit = int(source.get("Einstellung declinedFeeCents") or 0)
    check(all(cents(r["Ablehnungsgebuehr intern (EUR)"]) == fee_unit and r["Gebuehrenregel"] == "idle_in_shift" for r in declined),
          "M: every declined order carries the decline fee and its rule", f"{len(declined)} x {fee_unit} cents")
    _, revenue = table(ops, "N Interne Einnahmen")
    total = next((r for r in revenue if r["Posten"] == "Abgelehnte Auftraege gesamt"), None)
    check(total is not None and total["Anzahl"] == len(declined) and cents(total["Betrag (EUR)"]) == len(declined) * fee_unit,
          "N: the internal total equals M", f"{len(declined) * fee_unit / 100:.2f} EUR, not in the Fahrtenbuch")

    summary = {
        "report": report_path, "operational": ops_path, "verdict": verdict, "runId": source.get("Lauf-ID"),
        "rides": len(rides), "completed": len(done), "cancelled": len(cancelled), "declinedInternal": len(declined),
        "fareCents": fare, "tipCents": tips, "cancellationFeeCents": storno, "declinedFeeCentsInternal": len(declined) * fee_unit,
        "routeChecks": len(g_rows), "routeStatuses": dict(status), "legs": len(legs), "legStatuses": dict(leg_status),
        "shifts": len(accepted), "days": len(days), "zeroRideDays": len(zero),
    }
    return check, summary


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    check, summary = verify(argv[1], argv[2])
    print("\n".join(check.lines))
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print("VERIFY FAILED: " + "; ".join(check.failures) if check.failures else "VERIFY OK")
    return 1 if check.failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
